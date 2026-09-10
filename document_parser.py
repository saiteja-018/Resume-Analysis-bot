"""
Document parser for CareerMatch AI bot.
Extracts text from PDF and DOCX resume/JD files.
"""

import os
import logging
import pdfplumber
from docx import Document
from utils import clean_text

logger = logging.getLogger(__name__)


def extract_text_from_pdf(file_path: str) -> str:
    """
    Extract text from a PDF file using pdfplumber.
    Handles multi-page documents and complex layouts.
    """
    try:
        text_parts = []
        with pdfplumber.open(file_path) as pdf:
            if len(pdf.pages) == 0:
                return ""
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)

                # Also extract text from tables
                tables = page.extract_tables()
                for table in tables:
                    for row in table:
                        if row:
                            row_text = " | ".join(
                                cell.strip() if cell else "" for cell in row
                            )
                            if row_text.strip(" |"):
                                text_parts.append(row_text)

        full_text = "\n\n".join(text_parts)
        return clean_text(full_text)

    except Exception as e:
        logger.error(f"Failed to extract text from PDF: {e}")
        raise ValueError(f"Could not read the PDF file: {str(e)}")


def extract_text_from_docx(file_path: str) -> str:
    """
    Extract text from a DOCX file using python-docx.
    Extracts paragraphs and table content.
    """
    try:
        doc = Document(file_path)
        text_parts = []

        # Extract paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text.strip())

        # Extract table content
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip() for cell in row.cells if cell.text.strip()
                )
                if row_text:
                    text_parts.append(row_text)

        full_text = "\n".join(text_parts)
        return clean_text(full_text)

    except Exception as e:
        logger.error(f"Failed to extract text from DOCX: {e}")
        raise ValueError(f"Could not read the DOCX file: {str(e)}")


def extract_text(file_path: str, file_type: str) -> str:
    """
    Dispatch text extraction based on file type.

    Args:
        file_path: Path to the downloaded file.
        file_type: One of 'pdf' or 'docx'.

    Returns:
        Extracted text as a string.

    Raises:
        ValueError: If the file cannot be read or the type is unsupported.
    """
    if not os.path.exists(file_path):
        raise ValueError("File not found.")

    file_size = os.path.getsize(file_path)
    if file_size == 0:
        raise ValueError("The file is empty.")

    if file_type == "pdf":
        text = extract_text_from_pdf(file_path)
    elif file_type == "docx":
        text = extract_text_from_docx(file_path)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")

    if not text or len(text.strip()) < 20:
        raise ValueError(
            "Could not extract meaningful text from the file. "
            "It may be an image-based/scanned document or empty."
        )

    return text
